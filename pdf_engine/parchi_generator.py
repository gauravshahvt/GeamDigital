import os
import re
import subprocess
import tempfile
from io import BytesIO
from typing import List, Dict, Any, Optional, Tuple
import openpyxl

CHROME_PATHS = [
    r'C:\Program Files\Google\Chrome\Application\chrome.exe',
    r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe',
    r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',
    r'C:\Program Files\Microsoft\Edge\Application\msedge.exe'
]

def get_browser_executable() -> Optional[str]:
    """Finds installed Google Chrome or Microsoft Edge executable for PDF rendering."""
    for p in CHROME_PATHS:
        if os.path.exists(p):
            return p
    return None

def predict_gender(name: str, relative_name: str = '') -> str:
    """
    Infers gender (पुरूष / स्त्री) based on relationship and common Hindi name patterns.
    """
    name = (name or '').strip()
    rel = (relative_name or '').strip()

    # If 'पति' is in relationship or relative name -> स्त्री (Female)
    if 'पति' in rel or 'wife' in rel.lower():
        return 'स्त्री'

    female_keywords = [
        'देवी', 'बाई', 'कुमारी', 'बेगम', 'कंवर', 'कौर', 'सुगना', 'शांति', 'मंजू', 'मन्जू', 
        'रेखा', 'कमला', 'सीता', 'गीता', 'पुष्पा', 'सुशीला', 'अनीता', 'सुनीता', 'संगीता', 
        'माया', 'लीला', 'पूजा', 'आरती', 'राधा', 'ममता', 'शारदा', 'किरण', 'मोनिका', 
        'प्रियंका', 'नेनी', 'भागुती', 'मनभरी', 'काली', 'संतोष', 'ललिता', 'पिंकी', 'कौशल्या',
        'उर्मिला', 'मिस', 'विमला', 'भावना', 'आशा', 'कृष्णा', 'लक्ष्मी', 'पार्वती'
    ]
    for kw in female_keywords:
        if kw in name:
            return 'स्त्री'

    return 'पुरूष'

def clean_val_str(val: Any) -> str:
    """Converts cell values to clean strings, stripping .0 from floats."""
    if val is None:
        return ''
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return f"{val:.2f}"
    return str(val).strip()

ENGLISH_TO_HINDI_PANCHAYAT = {
    'ATOON': 'अटून', 'ATUN': 'अटून', 'AGOOCHA': 'आगूचा', 'AGUCHA': 'आगूचा',
    'AROLI': 'आरोली', 'BAGMALI': 'बागमाली', 'BAMANI': 'बामणी', 'BARANA': 'बराना',
    'DAULATGARH': 'दौलतगढ़', 'DANTDA': 'दांतड़ा', 'AMLIGARH': 'आमलीगढ़',
    'AMLIYA': 'आमलिया', 'ASIND': 'आसींद', 'ARJIYA': 'अर्जीया', 'BARDOD': 'बरडोद',
    'KASOR': 'कासोर', 'MOTRAS': 'मोतड़ास', 'ROOPAHELI': 'रूपाहेली', 'RUPAHELI': 'रूपाहेली',
    'BHILWARA': 'भीलवाड़ा'
}

def detect_panchayat_hindi(wb, file_path_or_bytes: Any = None, default_booth: str = '') -> str:
    """
    Detects the clean Hindi Gram Panchayat name from Excel workbook, sheet contents, booth addresses, and filename.
    Returns format e.g. 'ग्राम पंचायत अटून' or 'ग्राम पंचायत आगूचा'.
    """
    booths = []
    docs = []
    if default_booth:
        booths.append(str(default_booth).strip())

    try:
        for s_name in wb.sheetnames:
            ws = wb[s_name]
            for r in range(1, min(15, ws.max_row + 1)):
                for c in range(1, min(15, ws.max_column + 1)):
                    val = str(ws.cell(r, c).value or '').strip()
                    if not val:
                        continue
                    if any(k in val for k in ['रा.', 'विद्यालय', 'स्कूल', 'भवन', 'केंन्द', 'केंद्र', 'पंचायत']):
                        booths.append(val)
                    if '.pdf' in val.lower() or '.xlsx' in val.lower():
                        docs.append(val)
    except Exception:
        pass

    # 1. Check direct 'ग्राम पंचायत [नाम]' or 'ग्रामपंचायत [नाम]' pattern in any cell
    for b in booths:
        pm = re.search(r'(?:ग्राम\s*पंचायत|ग्रामपंचायत)\s*[:\-\s]*([A-Za-z\u0900-\u097F]+)', b)
        if pm:
            c = pm.group(1).strip()
            if len(c) > 2 and c not in ['नवीन', 'भवन', 'वार्ड', 'कमरा', 'कक्ष', 'प्राथमिक', 'विद्यालय']:
                return f'ग्राम पंचायत {c}'

    # 2. Check booth school address for Hindi village name after विद्यालय/वि./स्कूल/भवन/केंद्र
    stop_words = {'नवीन', 'पुराना', 'कमरा', 'क', 'कक्ष', 'प्राइमरी', 'उच्च', 'माध्यमिक', 'संख्या', 'न', 'नं', 'स', 'सं', 'क्षेत्र'}
    for b in booths:
        m = re.search(r'(?:विद्यालय|वि\.|स्कूल|भवन|केंन्द|केंद्र)\s*([A-Za-z\u0900-\u097F]+)', b)
        if m:
            c = m.group(1).strip()
            if c not in stop_words and len(c) >= 3:
                if re.match(r'^[A-Za-z]+$', c):
                    for eng, hin in ENGLISH_TO_HINDI_PANCHAYAT.items():
                        if eng.lower() == c.lower():
                            c = hin
                            break
                return f'ग्राम पंचायत {c}'

    # 3. Check ENGLISH_TO_HINDI_PANCHAYAT against docs and filename
    ref_text = ''
    if isinstance(file_path_or_bytes, (str, os.PathLike)):
        ref_text += ' ' + os.path.basename(file_path_or_bytes)
    ref_text += ' ' + ' '.join(docs)

    for eng, hin in ENGLISH_TO_HINDI_PANCHAYAT.items():
        if re.search(r'\b' + eng + r'\b', ref_text, re.I) or eng.lower() in ref_text.lower():
            return f'ग्राम पंचायत {hin}'

    # 4. Check any Devanagari in filename
    if isinstance(file_path_or_bytes, (str, os.PathLike)):
        base_fn = os.path.basename(file_path_or_bytes)
        clean_fn = re.sub(r'^(parchi|color_parchi|alpha|family)_[a-f0-9\-]+_', '', base_fn)
        clean_fn = re.sub(r'^[a-f0-9\-]{36}_', '', clean_fn)
        clean_fn = re.sub(r'\.xlsx?$|\.pdf$', '', clean_fn, flags=re.I)
        dev_m = re.search(r'[\u0900-\u097F]{3,}', clean_fn)
        if dev_m:
            c = dev_m.group(0)
            if c not in stop_words:
                return f'ग्राम पंचायत {c}'

    return 'ग्राम पंचायत'

def parse_excel_for_parchi(file_path_or_bytes: Any, sheet_name: Optional[str] = None) -> Dict[str, Any]:
    """
    Parses an uploaded Excel file for Voter Slip (Plain Matdata Parchi) generation.
    Supports single or multi-sheet Excel files.
    """
    if isinstance(file_path_or_bytes, (str, os.PathLike)):
        wb = openpyxl.load_workbook(file_path_or_bytes, data_only=True)
    else:
        wb = openpyxl.load_workbook(BytesIO(file_path_or_bytes), data_only=True)

    sheet_names = wb.sheetnames

    # Scan all workbook sheets (especially 'वार्ड_सारांश' / 'Summary') for JP, PS, Ward Map
    summary_jp_map = {}
    summary_default_jp = ""
    summary_default_ps = ""

    for s_name in sheet_names:
        if 'सारांश' in s_name or 'summary' in s_name.lower():
            s_ws = wb[s_name]
            sum_hdr_row = 1
            sum_col_map = {}
            for sr in range(1, min(6, s_ws.max_row + 1)):
                r_headers = [clean_val_str(s_ws.cell(row=sr, column=sc).value) for sc in range(1, min(15, s_ws.max_column + 1))]
                for idx, h in enumerate(r_headers, start=1):
                    h_clean = re.sub(r'[\s\.\/०\-_:;|]+', '', str(h).lower())
                    if any(k in h_clean for k in ['जिप', 'जिलापरिषद', 'jilaparishad']) or h_clean in ['jp', 'jpno']:
                        sum_col_map['jp'] = idx
                    elif any(k in h_clean for k in ['पंस', 'पंचायतसमिति', 'panchayatsamiti']) or h_clean in ['ps', 'psno', 'पंसमिति']:
                        sum_col_map['ps'] = idx
                    elif any(k in h_clean for k in ['वार्डनं', 'वार्डसंख्या', 'वार्डसं', 'wardno', 'ward_no']) or h_clean in ['वार्ड', 'ward']:
                        sum_col_map['ward'] = idx
                    elif any(k in h_clean for k in ['भागसंख्या', 'भागसं', 'partno']) or h_clean in ['भाग', 'part']:
                        sum_col_map['ward'] = idx
                if 'ward' in sum_col_map:
                    sum_hdr_row = sr
                    break

            for sr in range(sum_hdr_row + 1, min(40, s_ws.max_row + 1)):
                w_val = clean_val_str(s_ws.cell(row=sr, column=sum_col_map['ward']).value) if 'ward' in sum_col_map else ''
                w_match = re.search(r'\d+', w_val)
                if not w_match:
                    continue
                w_num = str(int(w_match.group(0)))

                s_jp = clean_val_str(s_ws.cell(row=sr, column=sum_col_map['jp']).value) if 'jp' in sum_col_map else ''
                s_ps = clean_val_str(s_ws.cell(row=sr, column=sum_col_map['ps']).value) if 'ps' in sum_col_map else ''

                # Sanitize: JP must not contain 'वार्ड' or 'ward'
                if 'वार्ड' in s_jp.lower() or 'ward' in s_jp.lower():
                    s_jp = ''
                if re.search(r'\.(pdf|xlsx?)$', s_ps, re.I) or 'वार्ड' in s_ps.lower() or 'ward' in s_ps.lower():
                    s_ps = ''

                if s_jp or s_ps:
                    summary_jp_map[w_num] = (s_jp, s_ps)
                if not summary_default_jp and s_jp and s_jp != '-':
                    summary_default_jp = s_jp
                if not summary_default_ps and s_ps and s_ps != '-':
                    summary_default_ps = s_ps

    # Select target sheets: prefer specific sheet, or 'समस्त_मतदाता_सूची', or all non-summary sheets
    target_sheets = []
    chosen_sheet = None
    if sheet_name and sheet_name in sheet_names:
        target_sheets = [sheet_name]
        chosen_sheet = sheet_name
    elif 'समस्त_मतदाता_सूची' in sheet_names:
        target_sheets = ['समस्त_मतदाता_सूची']
        chosen_sheet = 'समस्त_मतदाता_सूची'
    else:
        non_summary = [s for s in sheet_names if 'सारांश' not in s and 'summary' not in s.lower()]
        target_sheets = non_summary if non_summary else [sheet_names[0]]
        chosen_sheet = target_sheets[0]

    # Robust Column Mapping Helpers
    def match_col_field(text: str) -> Optional[str]:
        t_clean = re.sub(r'[\s\.\/०\-_:;|]+', '', str(text).lower())
        if any(k in t_clean for k in ['जिप', 'जिलापरिषद', 'jilaparishad', 'jilaparishadno']) or t_clean in ['jp', 'jpno']:
            return 'jila_parishad'
        if any(k in t_clean for k in ['पंस', 'पंचायतसमिति', 'panchayatsamiti', 'panchayatsamitino']) or t_clean in ['ps', 'psno', 'पंसमिति']:
            return 'panchayat_samiti'
        if any(k in t_clean for k in ['वार्डनं', 'वार्डसंख्या', 'वार्डसं', 'wardno', 'ward_no']) or t_clean in ['वार्ड', 'ward']:
            return 'part_no'
        if any(k in t_clean for k in ['भागसंख्या', 'भागसं', 'partno', 'part_no']) or t_clean in ['भाग', 'part']:
            return 'part_no'
        if any(k in t_clean for k in ['क्रमसंख्या', 'क्रमसं', 'क्रसं', 'क्रमांक', 'serial', 'srno', 'sno', 'slno']):
            return 'serial'
        if any(k in t_clean for k in ['मतदाताकानाम', 'निर्वाचककानाम', 'votername', 'fullname']) or (('नाम' in t_clean or 'name' in t_clean) and 'पिता' not in t_clean and 'पति' not in t_clean and 'बूथ' not in t_clean):
            return 'name'
        if any(k in t_clean for k in ['पिता/पति', 'पिताकानाम', 'पतिकानाम', 'पिता/पतिकानाम्', 'पितापतिकानाम्', 'संबंधीकानाम', 'fathername', 'husbandname', 'relativename']) or ('पिता' in t_clean or 'पति' in t_clean):
            return 'relative_name'
        if any(k in t_clean for k in ['आयु', 'उम्र', 'age', 'voterage']):
            return 'age'
        if any(k in t_clean for k in ['लिंग', 'जेंडर', 'gender', 'sex']):
            return 'gender'
        if any(k in t_clean for k in ['वोटरid', 'वोटरआईडी', 'epic', 'epicno', 'voterid', 'पहचानपत्रक्रमांक', 'पहचानपत्रसं', 'पहचानपत्र']):
            return 'epic'
        if any(k in t_clean for k in ['हाउसनंबर', 'हाउसनं', 'मकानसंख्या', 'मकाननं', 'मकानसं', 'मकाननंबर', 'houseno', 'housenumber', 'hno']):
            return 'house'
        if any(k in t_clean for k in ['बूथकापता', 'मतदानकेंद्रकापता', 'मतदानकेंद्रकानाम', 'मतदानकेंद्र', 'boothaddress', 'pollingstation', 'booth', 'बूथ']):
            return 'booth_address'
        if any(k in t_clean for k in ['status', 'स्थिति']):
            return 'status'
        return None

    # Extract Voters Data across target sheets
    voters = []
    default_booth = ""

    for s_item in target_sheets:
        ws = wb[s_item]
        s_digits = re.findall(r'\d+', s_item)
        sheet_ward_default = str(int(s_digits[0])) if s_digits else ''

        # Find Header Row (within first 5 rows)
        header_row_idx = 1
        for r in range(1, min(6, ws.max_row + 1)):
            row_vals = [clean_val_str(ws.cell(row=r, column=c).value).lower() for c in range(1, ws.max_column + 1)]
            if any(k in ' '.join(row_vals) for k in ['क्रम', 'नाम', 'serial', 'voter', 'epic', 'आयु', 'भाग']):
                header_row_idx = r
                break

        col_indices = {}
        for c in range(1, ws.max_column + 1):
            cell_val = clean_val_str(ws.cell(row=header_row_idx, column=c).value).strip()
            if not cell_val:
                continue
            field = match_col_field(cell_val)
            if field and field not in col_indices:
                col_indices[field] = c

        # Fallback by column position ONLY if column is not already mapped and contains genuine JP / PS headers
        if 'jila_parishad' not in col_indices and ws.max_column >= 1:
            c1_val = clean_val_str(ws.cell(row=header_row_idx, column=1).value).lower()
            c1_clean = re.sub(r'[\s\.\/०\-_:;|]+', '', c1_val)
            if any(k in c1_clean for k in ['जिप', 'जिलापरिषद', 'जिला']) or c1_clean in ['jp', 'jpno']:
                if 1 not in col_indices.values():
                    col_indices['jila_parishad'] = 1
        if 'panchayat_samiti' not in col_indices and ws.max_column >= 2:
            c2_val = clean_val_str(ws.cell(row=header_row_idx, column=2).value).lower()
            c2_clean = re.sub(r'[\s\.\/०\-_:;|]+', '', c2_val)
            if any(k in c2_clean for k in ['पंस', 'पंचायतसमिति', 'पंसमिति']) or c2_clean in ['ps', 'psno']:
                if 2 not in col_indices.values():
                    col_indices['panchayat_samiti'] = 2
        if 'part_no' not in col_indices and ws.max_column >= 3:
            c3_val = clean_val_str(ws.cell(row=header_row_idx, column=3).value).lower()
            c3_clean = re.sub(r'[\s\.\/०\-_:;|]+', '', c3_val)
            if any(k in c3_clean for k in ['वार्ड', 'ward', 'भाग', 'part']):
                if 3 not in col_indices.values():
                    col_indices['part_no'] = 3

        for r in range(header_row_idx + 1, ws.max_row + 1):
            row_cells = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            if not any(row_cells):
                continue

            def get_field(field_name: str) -> str:
                if field_name in col_indices:
                    return clean_val_str(ws.cell(row=r, column=col_indices[field_name]).value)
                return ''

            name = get_field('name')
            if not name:
                continue

            raw_serial = get_field('serial')
            try:
                serial_val = int(float(raw_serial)) if raw_serial else (len(voters) + 1)
            except (ValueError, TypeError):
                serial_val = len(voters) + 1

            part_no = get_field('part_no') or sheet_ward_default or '1'
            try:
                part_no = str(int(float(part_no)))
            except (ValueError, TypeError):
                pass

            epic = get_field('epic')
            relative_name = get_field('relative_name')
            raw_age = get_field('age')
            try:
                age_val = int(float(raw_age)) if raw_age else ''
            except (ValueError, TypeError):
                age_val = raw_age

            gender = get_field('gender')
            if not gender:
                gender = predict_gender(name, relative_name)
            elif gender.lower() in ['m', 'male', 'purush', 'पचरष']:
                gender = 'पुरूष'
            elif gender.lower() in ['f', 'female', 'mahila', 'stree', 'सल']:
                gender = 'स्त्री'

            house = get_field('house')
            booth = get_field('booth_address')
            if booth and not default_booth:
                default_booth = booth

            status = get_field('status') or 'Active'
            if 'del' in status.lower() or 'विलोपित' in status:
                status = 'Deleted'
            else:
                status = 'Active'

            v_jp = get_field('jila_parishad')
            v_ps = get_field('panchayat_samiti')

            # Sanitize: JP & PS must never be corrupted with ward number or file names
            if v_jp and ('वार्ड' in str(v_jp).lower() or 'ward' in str(v_jp).lower()):
                v_jp = ''
            if v_ps and (re.search(r'\.(pdf|xlsx?)$', str(v_ps), re.I) or 'वार्ड' in str(v_ps).lower() or 'ward' in str(v_ps).lower()):
                v_ps = ''

            # Fallback to summary map if cell was empty in voter row
            if not v_jp and part_no in summary_jp_map:
                v_jp = summary_jp_map[part_no][0]
            elif not v_jp and sheet_ward_default in summary_jp_map:
                v_jp = summary_jp_map[sheet_ward_default][0]
            elif not v_jp and summary_default_jp:
                v_jp = summary_default_jp

            if not v_ps and part_no in summary_jp_map:
                v_ps = summary_jp_map[part_no][1]
            elif not v_ps and sheet_ward_default in summary_jp_map:
                v_ps = summary_jp_map[sheet_ward_default][1]
            elif not v_ps and summary_default_ps:
                v_ps = summary_default_ps

            # Clean ward number
            m_w = re.search(r'\d+', str(part_no))
            clean_ward_no = m_w.group(0) if m_w else str(part_no).strip()

            voters.append({
                'serial': serial_val,
                'part_no': clean_ward_no,
                'ward': clean_ward_no,
                'jila_parishad': v_jp,
                'panchayat_samiti': v_ps,
                'name': name,
                'relative_name': relative_name,
                'age': age_val,
                'gender': gender,
                'epic': epic,
                'house': house if house else '00',
                'booth_address': booth,
                'status': status
            })

    # Auto-detect clean Hindi Gram Panchayat name from workbook metadata, booth addresses, and filename
    default_panchayat = detect_panchayat_hindi(wb, file_path_or_bytes, default_booth)

    default_jp = ""
    default_ps = ""
    default_part = "1"
    if voters:
        part_nos = sorted(list(set(v['part_no'] for v in voters if v['part_no'])))
        default_part = part_nos[0] if part_nos else "1"
        default_jp = next((str(v['jila_parishad']).strip() for v in voters if v.get('jila_parishad')), summary_default_jp)
        default_ps = next((str(v['panchayat_samiti']).strip() for v in voters if v.get('panchayat_samiti')), summary_default_ps)

    return {
        'sheet_names': sheet_names,
        'active_sheet': chosen_sheet,
        'voters': voters,
        'total_voters': len(voters),
        'active_count': len([v for v in voters if v['status'] == 'Active']),
        'deleted_count': len([v for v in voters if v['status'] == 'Deleted']),
        'default_panchayat_name': default_panchayat,
        'default_part_no': default_part,
        'default_booth_address': default_booth,
        'default_jila_parishad': default_jp,
        'default_panchayat_samiti': default_ps
    }

def generate_parchi_html(
    voters: List[Dict[str, Any]],
    panchayat_name: str = "ग्राम पंचायत",
    voting_time: str = "प्रातः 7 से सायं 6 बजे तक",
    filter_active_only: bool = True,
    layout_type: str = "12",
    show_page_number: bool = True,
    range_from: Optional[int] = None,
    range_to: Optional[int] = None,
    for_preview: bool = False,
    preview_page: int = 1,
    jila_parishad: Optional[str] = None,
    panchayat_samiti: Optional[str] = None
) -> Tuple[str, int, int]:
    """
    Renders HTML for Plain Matdata Parchi per A4 page.
    Layout Types:
      - '12': A4 - 12 (3 x 4 grid) [Default]
      - '16': A4 - 16 (4 x 4 grid)
      - '8' : A4 - 8  (2 x 4 grid)
    show_page_number: Toggle page number under each slip.
    """
    filtered = voters
    if filter_active_only:
        filtered = [v for v in voters if v.get('status') != 'Deleted']

    if range_from is not None or range_to is not None:
        r_from = range_from if range_from is not None else 1
        r_to = range_to if range_to is not None else 999999
        filtered = [v for v in filtered if r_from <= (v.get('serial') or 0) <= r_to]

    layout_type = str(layout_type).strip()
    if layout_type == "16":
        cols = 4
        rows = 4
        col_w = "49.5mm"
        cards_per_page = 16
        f_header = "11pt"
        f_voter_id = "7.5pt"
        f_data = "7.8pt"
        f_makan = "7.2pt"
        f_booth = "6.8pt"
        booth_height = "26px"
        f_time = "7pt"
        gap_w = "1mm"
    elif layout_type == "8":
        cols = 2
        rows = 4
        col_w = "99.5mm"
        cards_per_page = 8
        f_header = "15pt"
        f_voter_id = "10pt"
        f_data = "10pt"
        f_makan = "9.5pt"
        f_booth = "8.5pt"
        booth_height = "32px"
        f_time = "8.8pt"
        gap_w = "1.5mm"
    else:  # default "12" (3x4)
        cols = 3
        rows = 4
        col_w = "66mm"
        cards_per_page = 12
        f_header = "13.5pt"
        f_voter_id = "8.8pt"
        f_data = "8.8pt"
        f_makan = "8.2pt"
        f_booth = "7.5pt"
        booth_height = "28px"
        f_time = "8pt"
        gap_w = "1.2mm"

    card_height = "67.5mm" if show_page_number else "70.2mm"

    total_voters = len(filtered)
    total_pages = max(1, (total_voters + cards_per_page - 1) // cards_per_page)

    if for_preview:
        p_idx = max(1, min(preview_page, total_pages))
        page_chunks = [filtered[(p_idx - 1) * cards_per_page : p_idx * cards_per_page]]
        display_pages = [p_idx]
    else:
        page_chunks = [filtered[i * cards_per_page : (i + 1) * cards_per_page] for i in range(total_pages)]
        display_pages = list(range(1, total_pages + 1))

    html = f"""<!DOCTYPE html>
<html lang="hi">
<head>
<meta charset="utf-8">
<title>Geam Digital - Plain Matdata Parchi</title>
<style>
@page {{
    size: A4 portrait;
    margin: 4mm;
}}
* {{
    box-sizing: border-box;
    margin: 0;
    padding: 0;
}}
body {{
    font-family: 'Nirmala UI', 'Segoe UI', Arial, sans-serif;
    background: #fff;
    color: #000;
    -webkit-print-color-adjust: exact;
    print-color-adjust: exact;
}}
.sheet-page {{
    width: 202mm;
    height: 288mm;
    page-break-after: always;
    break-after: page;
    display: grid;
    grid-template-columns: repeat({cols}, {col_w});
    grid-template-rows: repeat(4, 70.8mm);
    gap: {gap_w};
    margin: 0 auto;
    overflow: hidden;
    position: relative;
}}
.parchi-cell {{
    display: flex;
    flex-direction: column;
    height: 70.8mm;
    justify-content: flex-start;
}}
.parchi-card {{
    border: 1px solid #111;
    padding: 2.5mm 2.8mm 2mm 2.8mm;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
    height: {card_height};
    background: #fff;
    overflow: hidden;
    box-sizing: border-box;
}}
.parchi-page-num {{
    text-align: right;
    font-size: 7.5pt;
    font-family: Arial, sans-serif;
    color: #000;
    padding-right: 1.5mm;
    line-height: 1.2;
    margin-top: 0.4mm;
}}
.header-group {{
    display: flex;
    flex-direction: column;
    gap: 1px;
}}
.header-title {{
    text-align: center;
    font-size: {f_header};
    font-weight: 700;
    line-height: 1.15;
    color: #000;
}}
.voter-id-row {{
    text-align: right;
    font-size: {f_voter_id};
    line-height: 1.2;
}}
.voter-id-val {{
    font-family: Arial, sans-serif;
    font-weight: 700;
    letter-spacing: 0.3px;
}}
.row-flex {{
    display: flex;
    justify-content: space-between;
    align-items: center;
    font-size: {f_data};
    line-height: 1.25;
}}
.row-item {{
    font-size: {f_data};
    line-height: 1.25;
}}
.bold-txt {{
    font-weight: 700;
}}
.makan-row {{
    text-align: right;
    font-size: {f_makan};
    line-height: 1.2;
}}
.booth-row {{
    font-size: {f_booth};
    line-height: 1.2;
    max-height: {booth_height};
    overflow: hidden;
    display: -webkit-box;
    -webkit-line-clamp: 2;
    -webkit-box-orient: vertical;
    text-overflow: ellipsis;
}}
.voting-time-box {{
    border: 1px solid #222;
    text-align: center;
    font-size: {f_time};
    font-weight: 700;
    padding: 2px 0;
    background: #fff;
    letter-spacing: 0.2px;
}}
.empty-card {{
    border: 1px dashed #ccc;
    background: #fafafa;
}}
@media print {{
    body {{
        margin: 0;
        padding: 0;
    }}
    .sheet-page {{
        margin: 0 auto;
    }}
}}
</style>
</head>
<body>
"""

    for page_num, chunk in zip(display_pages, page_chunks):
        html += f"""<div class="sheet-page" data-page="{page_num}">\n"""
        
        # Render cards per sheet
        for card_idx in range(cards_per_page):
            page_num_html = f"""<div class="parchi-page-num">{page_num}</div>""" if show_page_number else ""
            if card_idx < len(chunk):
                v = chunk[card_idx]
                serial = v.get('serial', '')
                ward = v.get('ward') or v.get('part_no', '')
                epic = v.get('epic', '')
                name = v.get('name', '')
                rel = v.get('relative_name', '')
                age = v.get('age', '')
                gender = v.get('gender', '')
                house = v.get('house', '00')
                booth = v.get('booth_address', '')
                jp = v.get('jila_parishad') or jila_parishad or ''
                ps = v.get('panchayat_samiti') or panchayat_samiti or ''

                # Gender & Age display (e.g. पुरूष 34 or स्त्री 28)
                gender_age_str = f"{gender} {age}".strip()

                card_html = f"""
    <div class="parchi-cell">
        <div class="parchi-card">
            <div class="header-group">
                <div class="header-title">{panchayat_name}</div>
                <div class="voter-id-row">Voter ID <span class="voter-id-val">{epic}</span></div>
            </div>
            <div class="row-flex">
                <div>क्रम सं. : <span class="bold-txt">{serial}</span></div>
                <div>वार्ड संख्या : <span class="bold-txt">{ward}</span></div>
            </div>
            <div class="row-flex">
                <div>जि. प. : <span class="bold-txt">{jp}</span></div>
                <div>पं. स. : <span class="bold-txt">{ps}</span></div>
            </div>
            <div class="row-flex">
                <div>नाम : <span class="bold-txt">{name}</span></div>
                <div class="bold-txt">{gender_age_str}</div>
            </div>
            <div class="row-flex">
                <div style="overflow: hidden; text-overflow: ellipsis; white-space: nowrap; max-width: 65%;">पिता/पति <span class="bold-txt">{rel}</span></div>
                <div style="white-space: nowrap;">मकानसं <span class="bold-txt">{house}</span></div>
            </div>
            <div class="booth-row">
                मतदान केंद्र - {booth}
            </div>
            <div class="voting-time-box">
                मतदान समय : {voting_time}
            </div>
        </div>
        {page_num_html}
    </div>
"""
            else:
                # Empty blank slot on final page if fewer than cards_per_page
                card_html = f"""
    <div class="parchi-cell">
        <div class="parchi-card empty-card"></div>
        {page_num_html}
    </div>
"""
            html += card_html

        html += """</div>\n"""

    html += """</body></html>"""
    return html, total_pages, total_voters

def generate_parchi_pdf(
    voters: List[Dict[str, Any]],
    panchayat_name: str = "ग्राम पंचायत",
    voting_time: str = "प्रातः 7 से सायं 6 बजे तक",
    filter_active_only: bool = True,
    layout_type: str = "12",
    show_page_number: bool = True,
    range_from: Optional[int] = None,
    range_to: Optional[int] = None,
    jila_parishad: Optional[str] = None,
    panchayat_samiti: Optional[str] = None
) -> bytes:
    """
    Generates a production-ready A4 PDF with Plain Matdata Parchi per sheet using headless Chrome/Edge.
    """
    html_content, total_pages, total_voters = generate_parchi_html(
        voters=voters,
        panchayat_name=panchayat_name,
        voting_time=voting_time,
        filter_active_only=filter_active_only,
        layout_type=layout_type,
        show_page_number=show_page_number,
        range_from=range_from,
        range_to=range_to,
        for_preview=False,
        jila_parishad=jila_parishad,
        panchayat_samiti=panchayat_samiti
    )

    browser_bin = get_browser_executable()
    if not browser_bin:
        raise RuntimeError("No headless Chrome or Edge browser found on system to generate PDF.")

    with tempfile.TemporaryDirectory() as tmpdir:
        html_file = os.path.join(tmpdir, "parchi.html")
        pdf_file = os.path.join(tmpdir, "parchi.pdf")

        with open(html_file, "w", encoding="utf-8") as f:
            f.write(html_content)

        cmd = [
            browser_bin,
            "--headless",
            "--disable-gpu",
            f"--print-to-pdf={pdf_file}",
            "--no-pdf-header-footer",
            html_file
        ]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        if res.returncode != 0 or not os.path.exists(pdf_file):
            raise RuntimeError(f"Browser PDF generation failed: {res.stderr}")

        with open(pdf_file, "rb") as f:
            pdf_bytes = f.read()

        return pdf_bytes
