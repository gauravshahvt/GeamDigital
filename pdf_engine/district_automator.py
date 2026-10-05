import os
import re
import time
from datetime import datetime
from typing import List, Dict, Any, Optional, Callable
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pdf_engine.voter_extractor import process_multiple_wards
from pdf_engine.excel_builder import create_excel_workbook

def scan_district_hierarchy(root_dir: str) -> Dict[str, Any]:
    """
    Scans a district root directory and automatically discovers:
    - Panchayat Samitis (parent folders)
    - Panchayats (folders containing ward PDFs)
    - Ward PDFs inside each Panchayat
    - Status of already generated Excel files (.xlsx)
    
    Resilient to:
    - Root / Samiti / Panchayat / *.pdf (3 levels)
    - Root / Panchayat / *.pdf (2 levels)
    - Mixed folder depths
    """
    root_dir = os.path.abspath(root_dir)
    if not os.path.exists(root_dir):
        raise FileNotFoundError(f"निर्दिष्ट डायरेक्टरी मौजूद नहीं है: {root_dir}")

    panchayats_found = []
    
    # Walk the directory tree to find folders containing PDF files
    for dirpath, dirnames, filenames in os.walk(root_dir):
        # Filter for PDF files
        pdf_files = [f for f in filenames if f.lower().endswith('.pdf') and not f.startswith('~$') and not f.startswith('.')]
        if not pdf_files:
            continue
            
        # Check for existing Excel file in this folder
        excel_files = [f for f in filenames if f.lower().endswith('.xlsx') and not f.startswith('~$') and not f.startswith('.')]
        existing_excel = excel_files[0] if excel_files else None
        
        # Calculate relative path to deduce Samiti and Panchayat names
        rel_path = os.path.relpath(dirpath, root_dir)
        path_parts = [p for p in rel_path.split(os.sep) if p and p != '.']
        
        if len(path_parts) == 0:
            # Root directory itself has PDFs
            samiti_name = os.path.basename(root_dir)
            panchayat_name = os.path.basename(root_dir)
        elif len(path_parts) == 1:
            # Root / Panchayat
            samiti_name = os.path.basename(root_dir)
            panchayat_name = path_parts[0]
        else:
            # Root / Samiti / Panchayat (or deeper)
            samiti_name = path_parts[-2]
            panchayat_name = path_parts[-1]
            
        full_pdf_paths = [os.path.join(dirpath, f) for f in pdf_files]
        
        panchayats_found.append({
            'samiti_name': samiti_name.strip(),
            'panchayat_name': panchayat_name.strip(),
            'folder_path': dirpath,
            'rel_path': rel_path,
            'pdf_files': pdf_files,
            'pdf_paths': full_pdf_paths,
            'pdf_count': len(pdf_files),
            'has_excel': bool(existing_excel),
            'existing_excel': existing_excel
        })
        
    # Group by Samiti
    samiti_map: Dict[str, List[Dict[str, Any]]] = {}
    for p in panchayats_found:
        s_name = p['samiti_name']
        if s_name not in samiti_map:
            samiti_map[s_name] = []
        samiti_map[s_name].append(p)
        
    total_pdfs = sum(p['pdf_count'] for p in panchayats_found)
    already_processed_count = sum(1 for p in panchayats_found if p['has_excel'])
    
    return {
        'root_dir': root_dir,
        'district_name': os.path.basename(root_dir),
        'total_samitis': len(samiti_map),
        'total_panchayats': len(panchayats_found),
        'total_pdfs': total_pdfs,
        'already_processed_count': already_processed_count,
        'pending_count': len(panchayats_found) - already_processed_count,
        'samitis': samiti_map,
        'panchayats': panchayats_found
    }


def generate_district_summary_excel(
    summary_data: List[Dict[str, Any]],
    output_path: str,
    district_name: str = "District"
) -> None:
    """
    Creates an Executive District-wide Excel report summarizing every Panchayat,
    Panchayat Samiti, Ward counts, and Voter figures.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    ws = wb.create_sheet(title="ज़िला_वोटर_मास्टर_सारांश")
    ws.views.sheetView[0].showGridLines = True
    
    # Colors
    navy_fill = PatternFill(start_color="0F2942", end_color="0F2942", fill_type="solid")
    gold_fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # 1. Main Banner
    ws.merge_cells("A1:J1")
    ws.row_dimensions[1].height = 40
    title_cell = ws["A1"]
    title_cell.value = f"GEAM DIGITAL - सम्पूर्ण ज़िला निर्वाचन सारांश: {district_name.upper()}"
    title_cell.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    title_cell.fill = navy_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 2. Subtitle / Timestamp
    ws.merge_cells("A2:J2")
    ws.row_dimensions[2].height = 22
    sub_cell = ws["A2"]
    sub_cell.value = f"रिपोर्ट जनरेट समय: {datetime.now().strftime('%d/%m/%Y, %I:%M %p')} | स्वतः निर्मित मल्टी-वार्ड एक्सेल सारांश"
    sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 3. Headers (Row 4)
    headers = [
        "क्र.सं.",
        "पंचायत समिति",
        "ग्राम पंचायत",
        "कुल वार्ड संख्या",
        "सक्रिय मतदाता",
        "विलोपित मतदाता",
        "कुल योग मतदाता",
        "एक्सेल फाइल",
        "स्थिति",
        "फ़ोल्डर पाथ"
    ]
    
    ws.row_dimensions[4].height = 28
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    # 4. Data Rows
    row_idx = 5
    tot_wards = 0
    tot_active = 0
    tot_del = 0
    tot_all = 0
    
    for idx, item in enumerate(summary_data, start=1):
        ws.row_dimensions[row_idx].height = 22
        is_even = (idx % 2 == 0)
        row_fill = zebra_fill if is_even else PatternFill(fill_type=None)
        
        wards_cnt = item.get('total_wards', 0)
        active_cnt = item.get('total_active_voters', 0)
        del_cnt = item.get('total_deleted_voters', 0)
        all_cnt = active_cnt + del_cnt
        
        tot_wards += wards_cnt
        tot_active += active_cnt
        tot_del += del_cnt
        tot_all += all_cnt
        
        values = [
            idx,
            item.get('samiti_name', ''),
            item.get('panchayat_name', ''),
            wards_cnt,
            active_cnt,
            del_cnt,
            all_cnt,
            item.get('excel_file', ''),
            item.get('status', 'सफल'),
            item.get('folder_path', '')
        ]
        
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            if row_fill.fill_type:
                cell.fill = row_fill
                
            if col_idx in (1, 4, 5, 6, 7):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if col_idx > 1 and isinstance(val, (int, float)):
                    cell.number_format = '#,##0'
            elif col_idx == 9:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if val == 'सफल':
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="047857")
                else:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
        row_idx += 1
        
    # 5. Grand Total Row
    ws.row_dimensions[row_idx].height = 28
    ws.cell(row=row_idx, column=1, value="").fill = gold_fill
    ws.cell(row=row_idx, column=2, value="").fill = gold_fill
    
    label_cell = ws.cell(row=row_idx, column=3, value="ज़िला कुल योग (Grand Total)")
    label_cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    label_cell.fill = gold_fill
    label_cell.alignment = Alignment(horizontal="center", vertical="center")
    label_cell.border = thin_border
    
    for col_idx, val in [(4, tot_wards), (5, tot_active), (6, tot_del), (7, tot_all)]:
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        cell.fill = gold_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        cell.number_format = '#,##0'
        
    for c in range(8, 11):
        cell = ws.cell(row=row_idx, column=c, value="")
        cell.fill = gold_fill
        cell.border = thin_border
        
    # Column Widths
    col_widths = {1: 8, 2: 24, 3: 26, 4: 16, 5: 16, 6: 16, 7: 18, 8: 34, 9: 14, 10: 45}
    for c, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
        
    wb.save(output_path)


def run_district_automation(
    root_dir: str,
    skip_existing: bool = True,
    theme: str = 'geam_digital',
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None
) -> Dict[str, Any]:
    """
    Automates the entire district voter PDF to Excel conversion:
    1. Scans root_dir and discovers all Panchayat Samitis and Panchayats.
    2. For each Panchayat, groups its ward PDFs.
    3. Runs multi-ward extraction.
    4. Generates multi-sheet Excel with Summary + All Wards.
    5. Saves the Excel directly in that Panchayat's folder.
    6. Generates a Master District Summary report in root_dir.
    """
    scan_info = scan_district_hierarchy(root_dir)
    panchayats = scan_info['panchayats']
    total_panchayats = len(panchayats)
    
    start_time = time.time()
    completed_panchayats = 0
    skipped_panchayats = 0
    failed_panchayats = 0
    total_voters = 0
    total_wards_processed = 0
    
    summary_records = []
    
    def emit(event_type: str, data: Dict[str, Any]):
        if progress_callback:
            try:
                progress_callback({
                    'type': event_type,
                    'timestamp': datetime.now().isoformat(),
                    'total_panchayats': total_panchayats,
                    'completed_panchayats': completed_panchayats,
                    'skipped_panchayats': skipped_panchayats,
                    'failed_panchayats': failed_panchayats,
                    'percent': round(((completed_panchayats + skipped_panchayats + failed_panchayats) / max(1, total_panchayats)) * 100, 1),
                    **data
                })
            except Exception:
                pass

    emit('start', {
        'message': f"ज़िला स्कैन पूरा हुआ: {scan_info['total_samitis']} पंचायत समितियां, {total_panchayats} पंचायतें, {scan_info['total_pdfs']} वार्ड पीडीएफ",
        'scan_info': scan_info
    })
    
    for idx, p_info in enumerate(panchayats, start=1):
        samiti = p_info['samiti_name']
        panchayat = p_info['panchayat_name']
        folder = p_info['folder_path']
        pdf_paths = p_info['pdf_paths']
        
        # Check if already generated
        if skip_existing and p_info['has_excel']:
            skipped_panchayats += 1
            emit('skip_panchayat', {
                'index': idx,
                'samiti': samiti,
                'panchayat': panchayat,
                'message': f"[{idx}/{total_panchayats}] {samiti} -> {panchayat}: पूर्व से मौजूद एक्सेल फाइल पाई गई ({p_info['existing_excel']}), स्किप की गई।"
            })
            summary_records.append({
                'samiti_name': samiti,
                'panchayat_name': panchayat,
                'total_wards': p_info['pdf_count'],
                'total_active_voters': 0,
                'total_deleted_voters': 0,
                'excel_file': p_info['existing_excel'],
                'status': 'स्किप (पूर्व से निर्मित)',
                'folder_path': folder
            })
            continue
            
        emit('start_panchayat', {
            'index': idx,
            'samiti': samiti,
            'panchayat': panchayat,
            'pdf_count': len(pdf_paths),
            'message': f"[{idx}/{total_panchayats}] प्रारंभ: {samiti} ➔ {panchayat} ({len(pdf_paths)} वार्ड्स)"
        })
        
        try:
            # Internal ward progress callback
            def ward_cb(info):
                emit('ward_progress', {
                    'index': idx,
                    'samiti': samiti,
                    'panchayat': panchayat,
                    'ward_idx': info.get('ward_idx', 1),
                    'total_wards': info.get('total_wards', len(pdf_paths)),
                    'message': info.get('status', '')
                })
                
            # Run multi-ward extractor
            res = process_multiple_wards(pdf_paths, progress_callback=ward_cb)
            
            # Generate Excel bytes
            excel_bytes = create_excel_workbook(
                tables=res['all_tables'],
                table_names=res['all_names'],
                theme=theme,
                mode='multi_sheet'
            )
            
            # Save file directly inside the Panchayat folder
            out_filename = res.get('suggested_filename')
            if not out_filename:
                out_filename = f"{panchayat} Ward 1 to {len(pdf_paths)}.xlsx"
            if not out_filename.endswith('.xlsx'):
                out_filename += '.xlsx'
                
            out_path = os.path.join(folder, out_filename)
            with open(out_path, 'wb') as f:
                f.write(excel_bytes)
                
            completed_panchayats += 1
            total_voters += res.get('total_active_voters', 0)
            total_wards_processed += len(pdf_paths)
            
            summary_records.append({
                'samiti_name': samiti,
                'panchayat_name': panchayat,
                'total_wards': len(pdf_paths),
                'total_active_voters': res.get('total_active_voters', 0),
                'total_deleted_voters': res.get('total_deleted_voters', 0),
                'excel_file': out_filename,
                'status': 'सफल',
                'folder_path': folder
            })
            
            emit('complete_panchayat', {
                'index': idx,
                'samiti': samiti,
                'panchayat': panchayat,
                'total_wards': len(pdf_paths),
                'active_voters': res.get('total_active_voters', 0),
                'deleted_voters': res.get('total_deleted_voters', 0),
                'excel_file': out_filename,
                'out_path': out_path,
                'message': f"[{idx}/{total_panchayats}] ✓ सफल: {panchayat} ({res.get('total_active_voters', 0)} मतदाता) ➔ {out_filename}"
            })
            
        except Exception as e:
            failed_panchayats += 1
            err_msg = str(e)
            emit('error_panchayat', {
                'index': idx,
                'samiti': samiti,
                'panchayat': panchayat,
                'error': err_msg,
                'message': f"[{idx}/{total_panchayats}] ❌ त्रुटि: {samiti} -> {panchayat}: {err_msg}"
            })
            summary_records.append({
                'samiti_name': samiti,
                'panchayat_name': panchayat,
                'total_wards': len(pdf_paths),
                'total_active_voters': 0,
                'total_deleted_voters': 0,
                'excel_file': '',
                'status': f'त्रुटि: {err_msg[:40]}',
                'folder_path': folder
            })
            
    # Generate Master District Summary Report
    elapsed_seconds = round(time.time() - start_time, 1)
    summary_path = os.path.join(root_dir, f"{scan_info['district_name']}_District_Voter_Summary.xlsx")
    try:
        generate_district_summary_excel(summary_records, summary_path, district_name=scan_info['district_name'])
    except Exception as e:
        print(f"Error generating district summary: {e}")
        
    emit('finish', {
        'message': f"संपूर्ण ज़िला ऑटोमेशन संपन्न! कुल {completed_panchayats} पंचायतें प्रोसेस हुईं ({total_voters} मतदाता)।",
        'summary_file': summary_path,
        'elapsed_seconds': elapsed_seconds
    })
    
    return {
        'success': True,
        'root_dir': root_dir,
        'district_name': scan_info['district_name'],
        'total_panchayats': total_panchayats,
        'completed_panchayats': completed_panchayats,
        'skipped_panchayats': skipped_panchayats,
        'failed_panchayats': failed_panchayats,
        'total_voters': total_voters,
        'total_wards_processed': total_wards_processed,
        'elapsed_seconds': elapsed_seconds,
        'summary_file': summary_path,
        'records': summary_records
    }
